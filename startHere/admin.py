from django.contrib import admin
from .models import (
    AKTFlashcard,
    AKTTextbookChapter,
    SCATextbookProgress,
    Category,
    Question,
    Option,
    QuestionImage,
    Record,
    Note,
    TextBook,
    Report,
)
from subscription.models import Package
from django import forms
from django.utils.html import format_html
from django_object_actions import DjangoObjectActions
from django.shortcuts import render
from django.contrib import messages
from django.http import HttpResponseRedirect, HttpResponse
from django.db import transaction
from django.urls import path, reverse
import pandas as pd
import openpyxl as xl
import re

class FileUploader(forms.Form):
    file = forms.FileField()


class AKTFlashcardUploadForm(forms.Form):
    file = forms.FileField(label='AKT flashcard workbook')
    packages = forms.ModelMultipleChoiceField(
        label='Make cards available in these packages',
        queryset=Package.objects.all(),
        help_text='The selected package assignments apply to every imported card.',
    )


class OptionInline(admin.TabularInline):
    model = Option
    extra = 4

class ImageInline(admin.TabularInline):
    model = QuestionImage
    extra = 4

@admin.register(Question)
class QuestionAdmin(admin.ModelAdmin):
    search_fields = ['content__icontains']
    autocomplete_fields = ["category"]
    inlines = [OptionInline,ImageInline]


@admin.register(TextBook)
class TextbookAdmin(admin.ModelAdmin):
    search_fields = ['content__icontains']
    autocomplete_fields = ["category"]


@admin.register(SCATextbookProgress)
class SCATextbookProgressAdmin(admin.ModelAdmin):
    list_display = ('user', 'textbook', 'is_complete', 'completed_at')
    list_filter = ('is_complete',)
    search_fields = ('user__username', 'textbook__title', 'textbook__category__title')
    autocomplete_fields = ('user', 'textbook')


class AKTTextbookUploadForm(forms.Form):
    file = forms.FileField(label='AKT textbook workbook')
    packages = forms.ModelMultipleChoiceField(
        label='Make chapters available in these packages',
        queryset=Package.objects.all(),
        help_text='Choose the package(s) once for this upload. These package assignments will be applied to every imported chapter.',
    )


@admin.register(AKTTextbookChapter)
class AKTTextbookChapterAdmin(admin.ModelAdmin):
    change_list_template = 'admin/startHere/akttextbookchapter/change_list.html'
    list_display = (
        'chapter_id', 'title', 'curriculum_area', 'topic', 'is_published', 'updated_at'
    )
    list_filter = ('curriculum_area', 'topic', 'is_published', 'package')
    search_fields = ('chapter_id', 'title', 'curriculum_area', 'topic')
    list_editable = ('is_published',)
    actions = ('publish_selected', 'unpublish_selected')
    readonly_fields = ('created_at', 'updated_at')

    IMPORT_COLUMNS = (
        'chapter_id', 'curriculum_area', 'topic', 'curriculum_outcomes',
        'title', 'order', 'body_html', 'sources'
    )

    def get_urls(self):
        custom_urls = [
            path(
                'bulk-upload/',
                self.admin_site.admin_view(self.bulk_upload_view),
                name='startHere_akttextbookchapter_bulk_upload',
            ),
            path(
                'download-template/',
                self.admin_site.admin_view(self.download_template),
                name='startHere_akttextbookchapter_download_template',
            ),
        ]
        return custom_urls + super().get_urls()

    @admin.action(description='Publish selected AKT chapters')
    def publish_selected(self, request, queryset):
        updated = queryset.update(is_published=True)
        self.message_user(request, f'{updated} AKT chapter(s) published.')

    @admin.action(description='Return selected AKT chapters to draft')
    def unpublish_selected(self, request, queryset):
        updated = queryset.update(is_published=False)
        self.message_user(request, f'{updated} AKT chapter(s) returned to draft.')

    def download_template(self, request):
        workbook = xl.Workbook()
        sheet = workbook.active
        sheet.title = 'AKT_Import'
        sheet.append(self.IMPORT_COLUMNS)
        sheet.freeze_panes = 'A2'
        for column, width in {
            'A': 34, 'B': 28, 'C': 30, 'D': 60, 'E': 40, 'F': 12, 'G': 80, 'H': 60,
        }.items():
            sheet.column_dimensions[column].width = width

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="GPReady_AKT_Textbook_Import_Template.xlsx"'
        workbook.save(response)
        return response

    def _render_upload(self, request, form, errors=None):
        context = {
            **self.admin_site.each_context(request),
            'opts': self.model._meta,
            'form': form,
            'errors': errors or [],
            'title': 'Bulk upload AKT textbook chapters',
        }
        return render(request, 'admin/startHere/akttextbookchapter/bulk_upload.html', context)

    def bulk_upload_view(self, request):
        form = AKTTextbookUploadForm(request.POST or None, request.FILES or None)
        if request.method != 'POST' or not form.is_valid():
            return self._render_upload(request, form)

        uploaded_file = form.cleaned_data['file']
        if not uploaded_file.name.lower().endswith('.xlsx'):
            form.add_error('file', 'Upload an .xlsx workbook.')
            return self._render_upload(request, form)

        errors = []
        parsed_rows = []
        seen_ids = set()
        try:
            workbook = xl.load_workbook(uploaded_file, read_only=True, data_only=True)
            if not workbook.sheetnames or workbook.sheetnames[0] != 'AKT_Import':
                errors.append('The first worksheet must be named AKT_Import.')
            else:
                sheet = workbook['AKT_Import']
                values = sheet.iter_rows(values_only=True)
                headers = next(values, None)
                normalised_headers = [str(value).strip().lower() if value is not None else '' for value in (headers or ())]
                missing = [column for column in self.IMPORT_COLUMNS if column not in normalised_headers]
                if missing:
                    errors.append('Missing required columns: ' + ', '.join(missing))
                else:
                    indexes = {column: normalised_headers.index(column) for column in self.IMPORT_COLUMNS}

                    def cell(row, column):
                        value = row[indexes[column]] if indexes[column] < len(row) else None
                        return '' if value is None else str(value).strip()

                    for row_number, row in enumerate(values, start=2):
                        if not row or not any(value not in (None, '') for value in row):
                            continue

                        chapter_id = cell(row, 'chapter_id').lower()
                        curriculum_area = cell(row, 'curriculum_area')
                        topic = cell(row, 'topic')
                        curriculum_outcomes = cell(row, 'curriculum_outcomes')
                        title = cell(row, 'title')
                        body_html = cell(row, 'body_html')
                        sources = cell(row, 'sources')
                        raw_order = row[indexes['order']] if indexes['order'] < len(row) else None

                        if not re.fullmatch(r'[a-z0-9]+(?:-[a-z0-9]+)*', chapter_id):
                            errors.append(f'Row {row_number}: chapter_id must be a lowercase, hyphenated ID.')
                        elif len(chapter_id) > 180:
                            errors.append(f'Row {row_number}: chapter_id must be 180 characters or fewer.')
                        elif chapter_id in seen_ids:
                            errors.append(f'Row {row_number}: duplicate chapter_id "{chapter_id}" in this workbook.')
                        else:
                            seen_ids.add(chapter_id)

                        for field_name, field_value in (
                            ('curriculum_area', curriculum_area),
                            ('topic', topic),
                            ('title', title),
                            ('body_html', body_html),
                        ):
                            if not field_value:
                                errors.append(f'Row {row_number}: {field_name} is required.')

                        for field_name, field_value, max_length in (
                            ('curriculum_area', curriculum_area, 180),
                            ('topic', topic, 220),
                            ('title', title, 255),
                        ):
                            if len(field_value) > max_length:
                                errors.append(f'Row {row_number}: {field_name} must be {max_length} characters or fewer.')

                        try:
                            position = int(raw_order) if raw_order not in (None, '') else 0
                            if isinstance(raw_order, float) and not raw_order.is_integer():
                                raise ValueError
                            if position < 0:
                                raise ValueError
                        except (TypeError, ValueError):
                            errors.append(f'Row {row_number}: order must be a whole number zero or greater.')
                            position = 0

                        parsed_rows.append({
                            'chapter_id': chapter_id,
                            'curriculum_area': curriculum_area,
                            'topic': topic,
                            'curriculum_outcomes': curriculum_outcomes,
                            'title': title,
                            'position': position,
                            'content': body_html,
                            'sources': sources,
                        })
        except Exception as exc:
            errors.append(f'Could not read the workbook: {exc}')

        if not parsed_rows:
            errors.append('The workbook contains no chapter rows.')

        if errors:
            return self._render_upload(request, form, errors)

        created_count = 0
        updated_count = 0
        with transaction.atomic():
            for chapter_data in parsed_rows:
                chapter, created = AKTTextbookChapter.objects.update_or_create(
                    chapter_id=chapter_data['chapter_id'],
                    defaults={
                        'curriculum_area': chapter_data['curriculum_area'],
                        'topic': chapter_data['topic'],
                        'curriculum_outcomes': chapter_data['curriculum_outcomes'],
                        'title': chapter_data['title'],
                        'position': chapter_data['position'],
                        'content': chapter_data['content'],
                        'sources': chapter_data['sources'],
                        # Imports always need a fresh review before the content is live.
                        'is_published': False,
                    },
                )
                chapter.package.set(form.cleaned_data['packages'])
                if created:
                    created_count += 1
                else:
                    updated_count += 1

        messages.success(
            request,
            f'Import complete: {created_count} chapter(s) created and {updated_count} updated. '
            'Imported chapters are drafts; review and publish them from the chapter list.',
        )
        return HttpResponseRedirect(reverse('admin:startHere_akttextbookchapter_changelist'))




@admin.register(Option)    
class OptionAdmin(admin.ModelAdmin):
    search_fields = ['content__icontains','content__istartswith']



@admin.register(Report)
class ReportAdmin(admin.ModelAdmin):
    list_display = ['user','question','comment','question_url']

    def question_url(self, report):
        url = f'/admin/startHere/question/{report.question.id}/change/'
        return format_html('<a href="{}">Edit question</a>', url)


admin.site.register(QuestionImage)
admin.site.register(Record)
admin.site.register(Note)

@admin.register(Category)
class CategoryAdmin(DjangoObjectActions, admin.ModelAdmin):
    autocomplete_fields = ["parent"]
    search_fields = ['title__icontains']

    def imports(modeladmin, request, queryset): 
        if request.method == "POST":
            file = request.FILES["file"]
        
            if not file.name.endswith('.xlsx'):
                messages.warning(request, 'The wrong file type was uploaded')
                return HttpResponseRedirect(request.path_info)
            
            
            data = pd.read_excel(file)
                        
            required_columns = ['category', 'question', 'a', 'b', 'c', 'd', 'e', 'f', 'correct','explanation']

        #     # Check if all required columns are present
            missing_columns = [col for col in required_columns if col not in data.columns]
            if missing_columns:
                messages.warning(request,f"The Excel file is missing the following required columns: {', '.join(missing_columns)}")
                return HttpResponseRedirect(request.path_info)
          
            with transaction.atomic():
                parent_category = Category.objects.get(id=queryset.id)                            
        #         # Create a list to hold the objects
                options_list = []
                         
        #         # Iterate through the rows in the Excel file
                for index, row in data.iterrows():
                    
                    if pd.isna(row['category']):
                        messages.warning(request,f"The field topic must not be empty")
                        return HttpResponseRedirect(request.path_info)
                    
                    if pd.isna(row['question']):
                        messages.warning(request,f"The field question must not be empty")
                        return HttpResponseRedirect(request.path_info)
                    
                    if pd.isna(row['correct']):
                        messages.warning(request,f"The field correct is empty for the question {row['question']}")
                        return HttpResponseRedirect(request.path_info)
                    else:
                        pattern = r'^[a-zA-Z]$'
                        if(bool(re.match(pattern, row['correct'].rstrip())) == False):
                            messages.warning(request,f"The field correct should contain only one alphabetic. {row['question']}")
                            return HttpResponseRedirect(request.path_info)
                            
                                        
                    category, _ = Category.objects.get_or_create(title=row['category'], parent=parent_category)
                    question, created = Question.objects.get_or_create(
                        category=category, 
                        content=row['question'],
                        explanation=row['explanation'] if pd.notna(row['explanation']) else None,
                    )

                    if created:
                        for option in ['a','b','c','d','e','f']:
                            if pd.notna(row[option]):
                                option = Option(
                                    question=question,
                                    content=row[option],
                                    is_correct=True if option == row['correct'] else False
                                )
                                options_list.append(option)
                            
                        
                        Option.objects.bulk_create(options_list)
                        options_list.clear()
                           
            messages.success(request, 'The file uploaded successfuly')
            return HttpResponseRedirect(request.path_info) 
            
        form = FileUploader()
        data = {"form": form}
        return render(request, "admin/upload.html",data)
        
        
    change_actions = ('imports', )












# @admin.register(Subject)
# class SubjectAdmin(DjangoObjectActions, admin.ModelAdmin):
#     inlines = [TopicInline]
#     list_display = ['title', 'subject_thumbnail']
#     search_fields = ['title__icontains']

#     def subject_thumbnail(self, instance):
#         return format_html(f'<img src={instance.thumbnail.url} class="thumbnail" />')

#     def exports(modeladmin, request, queryset):
#         subject = Subject.objects.get(id=queryset.id)
#         aphabetics = 'abcdefghijklmnopqrstuvwxyz'
          
#         data = []
#         for topic in subject.topics.all():
#             for question in topic.questions.all():
#                 options = question.options.all().values_list('content', flat=True)
#                 options_dict = {}
#                 # correct = ' '
#                 for index, option in enumerate(options):
#                     options_dict[aphabetics[index]] = option
#                 data.append({
#                     'topic': topic.title,
#                     'question': question.content,
#                     **options_dict,
#                     'correct' : ' '.join(aphabetics[index] for index, option in enumerate(question.options.all()) if option.is_correct),
#                     'year' : question.year_of_question,
#                     'type' : question.type,
#                     'explanation' : question.explanation
#                 })
        
            
#         # Create a DataFrame from the data
#         df = pd.DataFrame(data)

#         # Convert the DataFrame to an Excel file
#         with pd.ExcelWriter(f'{subject.title}.xlsx',  engine='openpyxl') as writer:
#             df.to_excel(writer, index=False)

#         # Return the Excel file as a response
#         with open(f'{subject.title}.xlsx', 'rb') as f:
#             response = HttpResponse(f.read(
#             ), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
#             response['Content-Disposition'] = f'attachment; filename={subject.title}.xlsx'
#             return response

    
#     def imports(modeladmin, request, queryset): 
#         if request.method == "POST":
#             file = request.FILES["file"]
        
#             if not file.name.endswith('.xlsx'):
#                 messages.warning(request, 'The wrong file type was uploaded')
#                 return HttpResponseRedirect(request.path_info)
            
            
#             data = pd.read_excel(file)
                        
#             required_columns = ['topic', 'question', 'a', 'b', 'c', 'd', 'e', 'f', 'correct','explanation']

#             # Check if all required columns are present
#             missing_columns = [col for col in required_columns if col not in data.columns]
#             if missing_columns:
#                 messages.warning(request,f"The Excel file is missing the following required columns: {', '.join(missing_columns)}")
#                 return HttpResponseRedirect(request.path_info)
          
#             with transaction.atomic():  
#                 subject = Subject.objects.get(id=queryset.id)
                            
#                 # Create a list to hold the objects
#                 options_list = []
                         
#                 # Iterate through the rows in the Excel file
#                 for index, row in data.iterrows():
                    
#                     if  pd.isna(row['topic']):
#                         messages.warning(request,f"The field topic must not be empty")
#                         return HttpResponseRedirect(request.path_info)
                    
#                     if  pd.isna(row['question']):
#                         messages.warning(request,f"The field question must not be empty")
#                         return HttpResponseRedirect(request.path_info)
                    
#                     if  pd.isna(row['correct']):
#                         messages.warning(request,f"The field correct is empty for the question {row['question']}")
#                         return HttpResponseRedirect(request.path_info)
#                     else:
#                         pattern = r'^[a-zA-Z]( [a-zA-Z])*$'
#                         if(bool(re.match(pattern, row['correct'].rstrip())) == False):
#                             messages.warning(request,f"The field correct should contain only alphabetic characters separated by spaces. {row['question']}")
#                             return HttpResponseRedirect(request.path_info)
                            
                                        
#                     topic, created = Topic.objects.get_or_create(title=row['topic'], subject=subject)
#                     question = Question.objects.create(
#                         topic=topic, 
#                         content=row['question'],
#                         explanation=row['explanation'] if pd.notna(row['explanation']) else None,
#                     )
                    
#                     for option in ['a','b','c','d','e','f']:
#                         if  pd.notna(row[option]):
                            
#                             option = Option(
#                                 question=question,
#                                 content=row[option],
#                                 is_correct=True if option in row['correct'].split(' ') else False
#                             )
#                             options_list.append(option)
                        
                    
#                     Option.objects.bulk_create(options_list)
#                     options_list.clear()
                           
#             messages.success(request, 'The file uploaded successfuly')
#             return HttpResponseRedirect(request.path_info) 
            
#         form = FileUploader()
#         data = {"form": form}
#         return render(request, "admin/upload.html",data)
        
        
#     change_actions = ('imports', )
#     # change_actions = ('imports', 'exports' ,)

    
#     class Media:
#         css = {
#             'all' : ['quizy/styles.css']
#         }


@admin.register(AKTFlashcard)
class AKTFlashcardAdmin(admin.ModelAdmin):
    change_list_template = 'admin/startHere/aktflashcard/change_list.html'
    list_display = ('card_id', 'deck', 'topic', 'difficulty', 'status', 'updated_at')
    list_filter = ('deck', 'difficulty', 'status', 'package')
    search_fields = ('card_id', 'deck', 'topic', 'front', 'answer', 'tags')
    list_editable = ('status',)
    actions = ('publish_selected', 'unpublish_selected')
    readonly_fields = ('created_at', 'updated_at')

    IMPORT_COLUMNS = (
        'card_id', 'deck', 'topic', 'front', 'clue', 'answer',
        'learning_point', 'source', 'tags', 'difficulty', 'status',
    )
    STATUS_VALUES = {'draft', 'reviewed', 'published'}
    DIFFICULTIES = {'foundation', 'intermediate', 'advanced'}

    def get_urls(self):
        custom_urls = [
            path(
                'bulk-upload/', self.admin_site.admin_view(self.bulk_upload_view),
                name='startHere_aktflashcard_bulk_upload',
            ),
            path(
                'download-template/', self.admin_site.admin_view(self.download_template),
                name='startHere_aktflashcard_download_template',
            ),
        ]
        return custom_urls + super().get_urls()

    @admin.action(description='Publish selected AKT flashcards')
    def publish_selected(self, request, queryset):
        updated = queryset.update(status='Published')
        self.message_user(request, f'{updated} flashcard(s) published.')

    @admin.action(description='Return selected AKT flashcards to draft')
    def unpublish_selected(self, request, queryset):
        updated = queryset.update(status='Draft')
        self.message_user(request, f'{updated} flashcard(s) returned to draft.')

    def download_template(self, request):
        workbook = xl.Workbook()
        sheet = workbook.active
        sheet.title = 'Flashcard_Import'
        sheet.append(self.IMPORT_COLUMNS)
        sheet.freeze_panes = 'A2'
        for column, width in {
            'A': 34, 'B': 24, 'C': 28, 'D': 58, 'E': 42, 'F': 58,
            'G': 48, 'H': 42, 'I': 28, 'J': 18, 'K': 16,
        }.items():
            sheet.column_dimensions[column].width = width
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="GPReady_AKT_Flashcard_Import_Template.xlsx"'
        workbook.save(response)
        return response

    def _render_flashcard_upload(self, request, form, errors=None):
        context = {
            **self.admin_site.each_context(request),
            'opts': self.model._meta,
            'form': form,
            'errors': errors or [],
            'title': 'Bulk upload AKT flashcards',
        }
        return render(request, 'admin/startHere/aktflashcard/bulk_upload.html', context)

    def bulk_upload_view(self, request):
        form = AKTFlashcardUploadForm(request.POST or None, request.FILES or None)
        if request.method != 'POST' or not form.is_valid():
            return self._render_flashcard_upload(request, form)

        uploaded_file = form.cleaned_data['file']
        if not uploaded_file.name.lower().endswith('.xlsx'):
            form.add_error('file', 'Upload an .xlsx workbook.')
            return self._render_flashcard_upload(request, form)

        errors = []
        parsed_rows = []
        seen_ids = set()
        try:
            workbook = xl.load_workbook(uploaded_file, read_only=True, data_only=True)
            if not workbook.sheetnames or workbook.sheetnames[0] != 'Flashcard_Import':
                errors.append('The first worksheet must be named Flashcard_Import.')
            else:
                rows = workbook['Flashcard_Import'].iter_rows(values_only=True)
                headers = next(rows, None)
                normalised = [str(value).strip().lower() if value is not None else '' for value in (headers or ())]
                missing = [column for column in self.IMPORT_COLUMNS if column not in normalised]
                if missing:
                    errors.append('Missing required columns: ' + ', '.join(missing))
                else:
                    indexes = {column: normalised.index(column) for column in self.IMPORT_COLUMNS}

                    def cell(row, column):
                        index = indexes[column]
                        value = row[index] if index < len(row) else None
                        return '' if value is None else str(value).strip()

                    for row_number, row in enumerate(rows, start=2):
                        if not row or not any(value not in (None, '') for value in row):
                            continue
                        card_id = cell(row, 'card_id').lower()
                        record = {column: cell(row, column) for column in self.IMPORT_COLUMNS}
                        if not re.fullmatch(r'[a-z0-9]+(?:-[a-z0-9]+)*', card_id):
                            errors.append(f'Row {row_number}: card_id must be a lowercase, hyphenated ID.')
                        elif len(card_id) > 180:
                            errors.append(f'Row {row_number}: card_id must be 180 characters or fewer.')
                        elif card_id in seen_ids:
                            errors.append(f'Row {row_number}: duplicate card_id "{card_id}" in this workbook.')
                        else:
                            seen_ids.add(card_id)

                        for field_name in ('deck', 'topic', 'front', 'answer'):
                            if not record[field_name]:
                                errors.append(f'Row {row_number}: {field_name} is required.')
                        for field_name, maximum in (('deck', 180), ('topic', 220), ('tags', 500)):
                            if len(record[field_name]) > maximum:
                                errors.append(f'Row {row_number}: {field_name} must be {maximum} characters or fewer.')
                        difficulty = record['difficulty'].title() or 'Intermediate'
                        if difficulty.lower() not in self.DIFFICULTIES:
                            errors.append(f'Row {row_number}: difficulty must be Foundation, Intermediate or Advanced.')
                        state = record['status'].lower() or 'draft'
                        if state not in self.STATUS_VALUES:
                            errors.append(f'Row {row_number}: status must be Draft, Reviewed or Published.')

                        parsed_rows.append({
                            'card_id': card_id,
                            'deck': record['deck'],
                            'topic': record['topic'],
                            'front': record['front'],
                            'clue': record['clue'],
                            'answer': record['answer'],
                            'learning_point': record['learning_point'],
                            'source': record['source'],
                            'tags': record['tags'],
                            'difficulty': difficulty,
                            # Draft and Reviewed content stays hidden until explicitly published.
                            'status': state.title(),
                        })
        except Exception as exc:
            errors.append(f'Could not read the workbook: {exc}')

        if not parsed_rows:
            errors.append('The workbook contains no flashcard rows.')
        if errors:
            return self._render_flashcard_upload(request, form, errors)

        created_count = updated_count = 0
        with transaction.atomic():
            for data in parsed_rows:
                card, created = AKTFlashcard.objects.update_or_create(
                    card_id=data.pop('card_id'), defaults=data,
                )
                card.package.set(form.cleaned_data['packages'])
                if created:
                    created_count += 1
                else:
                    updated_count += 1

        messages.success(
            request,
            f'Import complete: {created_count} flashcard(s) created and {updated_count} updated. '
            'Draft and Reviewed cards remain hidden until published.',
        )
        return HttpResponseRedirect(reverse('admin:startHere_aktflashcard_changelist'))

